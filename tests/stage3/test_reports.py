"""AC-15 and FR-RP-01…04: a one-country report from the admin, XLSX + methodology, no cell < 30."""

from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from payintel.api.deps import AppState
from payintel.core.clock import FixedClock
from payintel.core.models.base import ProviderRole
from payintel.reports import aggregates
from payintel.reports import service as reports
from payintel.reports.wilson import wilson
from tests.stage3.conftest import StoreSpec, World, login, make_store, staff_principal

pytestmark = pytest.mark.integration


def _seed_market(db_session: Session, clock: FixedClock) -> None:
    """DE: 35 shopify stores (adyen 30, stripe 5 → merged), 4 magento2 stores (cell withheld)."""
    for i in range(35):
        provider = "adyen" if i < 30 else "stripe"
        make_store(
            db_session,
            StoreSpec(
                f"de-shop-{i}.de",
                platform="shopify",
                providers=((provider, ProviderRole.GATEWAY),)
                + ((("klarna", ProviderRole.BNPL),) if i % 2 == 0 else ()),
                methods=("visa", "paypal") if i % 3 else ("visa",),
                rank=10_000 + i,
            ),
            clock=clock,
        )
    for i in range(4):
        make_store(
            db_session,
            StoreSpec(f"de-mag-{i}.de", platform="magento2", rank=20_000 + i),
            clock=clock,
        )


def test_report_suppresses_small_cells(
    db_session: Session, world: World, app_state: AppState, fixed_clock: FixedClock
) -> None:
    _seed_market(db_session, fixed_clock)
    spec = aggregates.ReportSpec(countries=("DE",), min_cell=30)
    report = aggregates.build(db_session, spec, now=fixed_clock.now())
    assert report.total_stores == 3 + 35 + 4  # world DE stores + seeded
    published = {k: v for k, v in report.cells.items() if v >= 30}
    assert published == {"DE|shopify": 37}  # alpha-shop + gamma-market are shopify too
    assert report.suppressed_cells >= 2  # woocommerce (1), magento2 (4)
    assert report.min_published() >= 30
    psp = {r.key: r.stores for r in report.psp_share if r.group == "DE|shopify"}
    assert psp["adyen"] == 31 and "stripe" not in psp and psp["other"] >= 5  # alpha-shop: adyen
    for r in report.psp_share + report.method_share + report.bnpl_share:
        assert r.group == "DE|shopify" and (r.key == "other" or r.stores >= 30)
        lo, hi = r.ci
        assert 0 <= lo <= r.share <= hi <= 1
    assert (
        "DE|shopify" in report.providers_per_store and report.providers_per_store["DE|shopify"] > 1
    )
    assert all(n >= 30 for _, _, n in report.monthly)


def test_admin_builds_xlsx_with_methodology(
    client: TestClient,
    db_session: Session,
    world: World,
    app_state: AppState,
    fixed_clock: FixedClock,
) -> None:
    _seed_market(db_session, fixed_clock)
    csrf = login(client, world.staff_analyst_email, state=app_state, session=db_session)
    r = client.post(
        "/admin/reports",
        data={
            "csrf": csrf,
            "countries": "de",
            "metrics": list(aggregates.METRICS),
            "org_id": str(world.org.id),
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text
    job = reports.recent(db_session, limit=1)[0]
    assert (
        job.status.value == "done" and job.xlsx_key and job.csv_key and job.org_id == world.org.id
    )
    assert job.summary["min_published"] >= 30 and job.summary["cells_withheld"] >= 2
    x = client.get(f"/admin/reports/{job.id}/xlsx")
    assert x.status_code == 200 and x.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Summary", "PSP share", "Methodology"} <= set(wb.sheetnames)
    meth = "\n".join(str(c.value) for row in wb["Methodology"].iter_rows() for c in row if c.value)
    for needle in ("Sample", "Coverage", "Known limitations", "Wilson", "Cell suppression"):
        assert needle in meth, needle
    for row in wb["PSP share"].iter_rows(min_row=2, values_only=True):
        assert row[1] == "other" or int(str(row[2])) >= 30
    z = client.get(f"/admin/reports/{job.id}/csv")
    assert z.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert any(n.startswith("psp_share") for n in names) and any("methodology" in n for n in names)
    # the client organisation sees and downloads its report from the portal
    client.cookies.clear()
    login(client, world.org_admin_email, state=app_state, session=db_session)
    page = client.get("/portal/reports").text
    assert "XLSX" in page and str(job.id) in page
    assert client.get(f"/portal/reports/{job.id}/xlsx").status_code == 200
    assert client.get(f"/portal/reports/{job.id}/pdf").status_code == 404
    # another organisation cannot
    from payintel.core.models.base import Role
    from tests.stage3.conftest import new_user

    new_user(db_session, "other@other.example", Role.ORG_ADMIN, world.other_org, clock=fixed_clock)
    client.cookies.clear()
    login(client, "other@other.example", state=app_state, session=db_session)
    assert client.get(f"/portal/reports/{job.id}/xlsx").status_code == 404


def test_report_requires_analyst_role(
    client: TestClient, db_session: Session, world: World, app_state: AppState
) -> None:
    csrf = login(client, world.staff_support_email, state=app_state, session=db_session)
    assert client.get("/admin/reports").status_code == 200
    assert client.post("/admin/reports", data={"csrf": csrf, "countries": "DE"}).status_code == 403
    from payintel.core.models.base import Role
    from payintel.entitlements.check import EntitlementDenied
    from payintel.entitlements.model import Principal

    with pytest.raises(EntitlementDenied):
        reports.build_report(
            db_session,
            aggregates.ReportSpec(countries=("DE",)),
            principal=Principal(kind="staff", org_id=None, role=Role.STAFF_SUPPORT, email="s@x"),
            store=None,
            settings=app_state.settings,
            clock=app_state.clock,
        )
    _ = staff_principal


def test_wilson_interval() -> None:
    assert wilson(0, 0) == (0.0, 0.0)
    lo, hi = wilson(30, 100)
    assert 0.21 < lo < 0.30 < hi < 0.40
    lo, hi = wilson(100, 100)
    assert lo > 0.96 and hi == pytest.approx(1.0)
    assert wilson(5, 3) == wilson(3, 3)
