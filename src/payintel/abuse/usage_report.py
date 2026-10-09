"""Quarterly usage report per organisation (FR-AB-05).

Built from the usage journal, the export jobs, the alert deliveries and the
incidents of the quarter: an XLSX (Summary, By endpoint, By day, Exports,
Incidents) stored in the exports bucket under `usage/<org>/<period>.xlsx`,
plus a JSON summary kept in `usage_report` for the portal and the admin.
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from payintel.core.models.abuse import AbuseIncident, UsageReport
from payintel.core.models.alerts import AlertRule, Delivery
from payintel.core.models.audit import UsageLog
from payintel.core.models.base import OrgStatus
from payintel.core.models.exports import ExportJob
from payintel.core.models.orgs import Organization
from payintel.core.s3 import ObjectStore


@dataclass(frozen=True)
class Quarter:
    year: int
    q: int

    @property
    def period(self) -> str:
        return f"{self.year}-Q{self.q}"

    @property
    def start(self) -> datetime:
        return datetime(self.year, 3 * (self.q - 1) + 1, 1, tzinfo=UTC)

    @property
    def end(self) -> datetime:
        if self.q == 4:
            return datetime(self.year + 1, 1, 1, tzinfo=UTC)
        return datetime(self.year, 3 * self.q + 1, 1, tzinfo=UTC)

    @classmethod
    def of(cls, d: date) -> Quarter:
        return cls(d.year, (d.month - 1) // 3 + 1)

    def previous(self) -> Quarter:
        return Quarter(self.year - 1, 4) if self.q == 1 else Quarter(self.year, self.q - 1)


@dataclass
class UsageSummary:
    period: str
    requests: int = 0
    records: int = 0
    errors: int = 0
    unique_ips: int = 0
    by_endpoint: list[tuple[str, int, int]] = field(default_factory=list)
    by_day: list[tuple[str, int, int]] = field(default_factory=list)
    exports: list[tuple[str, str, str, int]] = field(default_factory=list)
    alerts_delivered: int = 0
    incidents: list[tuple[str, str, str, str]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "period": self.period,
            "requests": self.requests,
            "records": self.records,
            "errors": self.errors,
            "unique_ips": self.unique_ips,
            "endpoints": len(self.by_endpoint),
            "exports": len(self.exports),
            "alerts_delivered": self.alerts_delivered,
            "incidents": len(self.incidents),
        }


def summarise(session: Session, org_id: uuid.UUID, quarter: Quarter) -> UsageSummary:
    start, end = quarter.start, quarter.end
    s = UsageSummary(quarter.period)
    base = (UsageLog.org_id == org_id, UsageLog.ts >= start, UsageLog.ts < end)
    s.requests, s.records, s.unique_ips = (
        int(x or 0)
        for x in session.execute(
            select(
                func.count(UsageLog.id),
                func.coalesce(func.sum(UsageLog.records), 0),
                func.count(func.distinct(UsageLog.ip)),
            ).where(*base)
        ).one()
    )
    s.errors = int(
        session.execute(
            select(func.count(UsageLog.id)).where(*base, UsageLog.params["status"].astext >= "400")
        ).scalar_one()
    )
    s.by_endpoint = [
        (str(e), int(n), int(r or 0))
        for e, n, r in session.execute(
            select(UsageLog.endpoint, func.count(UsageLog.id), func.sum(UsageLog.records))
            .where(*base)
            .group_by(UsageLog.endpoint)
            .order_by(func.count(UsageLog.id).desc())
        )
    ]
    day = func.date_trunc("day", UsageLog.ts)
    s.by_day = [
        (d.date().isoformat(), int(n), int(r or 0))
        for d, n, r in session.execute(
            select(day, func.count(UsageLog.id), func.sum(UsageLog.records))
            .where(*base)
            .group_by(day)
            .order_by(day)
        )
    ]
    s.exports = [
        (str(j.id), j.created_at.date().isoformat(), j.status.value, int(j.rows or 0))
        for j in session.execute(
            select(ExportJob)
            .where(
                ExportJob.org_id == org_id,
                ExportJob.created_at >= start,
                ExportJob.created_at < end,
            )
            .order_by(ExportJob.created_at)
        ).scalars()
    ]
    s.alerts_delivered = int(
        session.execute(
            select(func.count(Delivery.id))
            .join(AlertRule, AlertRule.id == Delivery.alert_rule_id)
            .where(
                AlertRule.org_id == org_id,
                Delivery.delivered_at.is_not(None),
                Delivery.delivered_at >= start,
                Delivery.delivered_at < end,
            )
        ).scalar_one()
    )
    s.incidents = [
        (i.created_at.date().isoformat(), i.detector, i.severity, i.status)
        for i in session.execute(
            select(AbuseIncident)
            .where(
                AbuseIncident.org_id == org_id,
                AbuseIncident.created_at >= start,
                AbuseIncident.created_at < end,
            )
            .order_by(AbuseIncident.created_at)
        ).scalars()
    ]
    return s


def write_xlsx(org: Organization, s: UsageSummary, *, generated_at: datetime) -> bytes:
    wb = Workbook()
    ws = wb.create_sheet("Summary", 0)
    wb.remove(wb.worksheets[1])
    ws.append([f"PayIntel usage report {s.period}"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    for k, v in (
        ("organisation", org.legal_name),
        ("organisation_id", str(org.id)),
        ("generated_at", generated_at.isoformat()),
        ("requests", s.requests),
        ("records", s.records),
        ("errors_4xx_5xx", s.errors),
        ("unique_ips", s.unique_ips),
        ("exports", len(s.exports)),
        ("alerts_delivered", s.alerts_delivered),
        ("incidents", len(s.incidents)),
    ):
        ws.append([k, v])
    sheets: list[tuple[str, list[str], list[tuple[Any, ...]]]] = [
        ("By endpoint", ["endpoint", "requests", "records"], list(s.by_endpoint)),
        ("By day", ["day", "requests", "records"], list(s.by_day)),
        ("Exports", ["export_id", "created", "status", "rows"], list(s.exports)),
        ("Incidents", ["created", "detector", "severity", "status"], list(s.incidents)),
    ]
    for title, header, rows in sheets:
        sheet = wb.create_sheet(title)
        sheet.append(header)
        for row in rows:
            sheet.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build(
    session: Session,
    org: Organization,
    quarter: Quarter,
    *,
    store: ObjectStore | None,
    bucket: str,
    now: datetime,
) -> UsageReport:
    s = summarise(session, org.id, quarter)
    key: str | None = None
    if store is not None:
        key = f"usage/{org.id}/{quarter.period}.xlsx"
        store.put_bytes(
            bucket,
            key,
            write_xlsx(org, s, generated_at=now),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    row = session.execute(
        select(UsageReport).where(
            UsageReport.org_id == org.id, UsageReport.period == quarter.period
        )
    ).scalar_one_or_none()
    if row is None:
        row = UsageReport(
            org_id=org.id,
            period=quarter.period,
            file_key=key,
            summary=s.as_dict(),
            generated_at=now,
        )
        session.add(row)
    else:
        row.file_key, row.summary, row.generated_at = key, s.as_dict(), now
    session.flush()
    return row


def build_all(
    session: Session,
    quarter: Quarter,
    *,
    store: ObjectStore | None,
    bucket: str,
    now: datetime,
    org_ids: list[uuid.UUID] | None = None,
) -> list[UsageReport]:
    q = select(Organization)
    if org_ids:
        q = q.where(Organization.id.in_(org_ids))
    else:
        q = q.where(Organization.status.in_([OrgStatus.ACTIVE, OrgStatus.SUSPENDED]))
    return [
        build(session, org, quarter, store=store, bucket=bucket, now=now)
        for org in session.execute(q.order_by(Organization.legal_name)).scalars()
    ]


def reports_of(session: Session, org_id: uuid.UUID) -> list[UsageReport]:
    return list(
        session.execute(
            select(UsageReport)
            .where(UsageReport.org_id == org_id)
            .order_by(UsageReport.period.desc())
        ).scalars()
    )
